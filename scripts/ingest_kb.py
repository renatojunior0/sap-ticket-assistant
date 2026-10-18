"""Converte data/raw_docs/ em data/kb_chunks.jsonl.
Uso: python scripts/ingest_kb.py [--ocr] [--include-all]
  --ocr          aplica Tesseract em PDFs sem camada de texto (requer tesseract-ocr instalado)
  --include-all  ignora a lista de exclusão (materiais longos/licenciados)
"""
import argparse
import io
import json
import re
from pathlib import Path

import pymupdf
import openpyxl
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw_docs"
OUTPUT = ROOT / "data" / "kb_chunks.jsonl"
CHUNK_CHARS, OVERLAP, MIN_PAGE_TEXT = 1200, 200, 40
EXCLUDE_PATTERNS = ("TS450", "TS452", "S99513", "SAP WM", "Transações sap")  # diluem o retriever; planilha vira ontologia


def pdf_pages(path: Path, ocr: bool) -> tuple[list[str], str]:
    document = pymupdf.open(path)
    pages = [page.get_text() for page in document]
    if sum(len(p.strip()) for p in pages) >= MIN_PAGE_TEXT * len(pages) or not ocr:
        return pages, "texto"
    import pytesseract
    from PIL import Image

    pages = [
        pytesseract.image_to_string(Image.open(io.BytesIO(page.get_pixmap(dpi=200).tobytes("png"))), lang="por+eng")
        for page in document
    ]
    return pages, "ocr"


def docx_pages(path: Path, _: bool) -> tuple[list[str], str]:
    document = Document(path)
    paragraphs = [p.text for p in document.paragraphs]
    tables = [" | ".join(c.text for c in row.cells) for t in document.tables for row in t.rows]
    return ["\n".join(paragraphs + tables)], "texto"


def xlsx_pages(path: Path, _: bool) -> tuple[list[str], str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = (" | ".join(str(c) for c in row if c is not None) for s in workbook for row in s.iter_rows(values_only=True))
    return ["\n".join(rows)], "texto"


READERS = {".pdf": pdf_pages, ".docx": docx_pages, ".xlsx": xlsx_pages}


def split(text: str) -> list[str]:
    text = re.sub(r"\s+", " ", text).strip()
    step = CHUNK_CHARS - OVERLAP
    return [text[i : i + CHUNK_CHARS] for i in range(0, len(text), step)] if text else []


def chunk_document(pages: list[str]) -> list[dict]:
    """Agrupa páginas até CHUNK_CHARS preservando o intervalo de páginas para citação."""
    chunks, buffer, start = [], "", 1
    for number, page in enumerate(pages, start=1):
        page = page.strip()
        if not page:
            continue
        if buffer and len(buffer) + len(page) > CHUNK_CHARS:
            chunks += [{"text": t, "pages": f"{start}-{number - 1}" if number - 1 > start else str(start)} for t in split(buffer)]
            buffer, start = "", number
        buffer += " " + page
    if buffer:
        chunks += [{"text": t, "pages": f"{start}-{len(pages)}" if len(pages) > start else str(start)} for t in split(buffer)]
    return chunks


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--ocr", action="store_true")
    parser.add_argument("--include-all", action="store_true")
    args = parser.parse_args()

    total, excluded, empty = 0, [], []
    with OUTPUT.open("w", encoding="utf-8") as out:
        for path in sorted(RAW_DIR.iterdir()):
            reader = READERS.get(path.suffix.lower())
            if not reader:
                continue
            if not args.include_all and any(p.lower() in path.name.lower() for p in EXCLUDE_PATTERNS):
                excluded.append(path.name)
                continue
            try:
                pages, mode = reader(path, args.ocr)
                chunks = chunk_document(pages)
            except Exception as exc:
                print(f"falha {path.name}: {type(exc).__name__}: {str(exc)[:120]}")
                continue
            if not chunks:
                empty.append(path.name)
            for index, chunk in enumerate(chunks):
                out.write(json.dumps({"id": f"{path.stem}#{index}", "source": path.name, "pages": chunk["pages"],
                                      "mode": mode, "text": chunk["text"]}, ensure_ascii=False) + "\n")
            total += len(chunks)
            print(f"{path.name}: {len(chunks)} chunks ({mode})")
    print(f"\ntotal={total} -> {OUTPUT.relative_to(ROOT)}")
    print(f"excluídos ({len(excluded)}): {excluded}")
    print(f"sem texto ({len(empty)}): {empty}  {'<- rode com --ocr' if empty and not args.ocr else ''}")


if __name__ == "__main__":
    main()